"""9.3 «Export»: fills the Πρότυπο 1 notice (.docx) from an announcement.

`docx_templates/notice_template.docx` is the Πρότυπο 1 as issued (sample: Λευκωσία, 28
Σεπτεμβρίου 2026, 2 μήνες, 30/11/2026). Only what the announcement records is replaced:
the notice date, the district, the communities with their number of plots, the submission
period in months and the last day of submission. The rest of the text is left as it is.
"""

import copy
import unicodedata
from collections import Counter
from io import BytesIO
from pathlib import Path

from docx import Document

from core.docx_merge import replace_in_paragraph

TEMPLATE_PATH = Path(__file__).with_name("docx_templates") / "notice_template.docx"

# The values of the sample notice, replaced as they stand in the template.
SAMPLE_DATE = "28 Σεπτεμβρίου, 2026"
SAMPLE_DISTRICT = "Λευκωσίας"
SAMPLE_DISTRICT_UPPER = "ΛΕΥΚΩΣΙΑΣ"
SAMPLE_PERIOD = "δύο (2) μηνών"
SAMPLE_END_DATE = "30/11/2026"

MONTHS_GENITIVE = (
    "Ιανουαρίου", "Φεβρουαρίου", "Μαρτίου", "Απριλίου", "Μαΐου", "Ιουνίου",
    "Ιουλίου", "Αυγούστου", "Σεπτεμβρίου", "Οκτωβρίου", "Νοεμβρίου", "Δεκεμβρίου",
)  # fmt: skip

PERIOD_IN_MONTHS = {
    1: "ενός (1) μηνός",
    2: "δύο (2) μηνών",
    3: "τριών (3) μηνών",
    4: "τεσσάρων (4) μηνών",
    5: "πέντε (5) μηνών",
    6: "έξι (6) μηνών",
    7: "επτά (7) μηνών",
    8: "οκτώ (8) μηνών",
    9: "εννέα (9) μηνών",
    10: "δέκα (10) μηνών",
    11: "έντεκα (11) μηνών",
    12: "δώδεκα (12) μηνών",
}

# Community.district holds the official upper-case genitive («ΛΕΥΚΩΣΙΑΣ»); the notice also
# needs it in running text, with accents.
DISTRICT_GENITIVE = {
    "ΛΕΥΚΩΣΙΑΣ": "Λευκωσίας",
    "ΛΕΜΕΣΟΥ": "Λεμεσού",
    "ΛΑΡΝΑΚΑΣ": "Λάρνακας",
    "ΠΑΦΟΥ": "Πάφου",
    "ΑΜΜΟΧΩΣΤΟΥ": "Αμμοχώστου",
    "ΚΕΡΥΝΕΙΑΣ": "Κερύνειας",
}


class NoticeExportError(Exception):
    """The announcement cannot be turned into a notice; the message is shown to the user."""


def _strip_accents(text: str) -> str:
    decomposed = unicodedata.normalize("NFD", text)
    return unicodedata.normalize("NFC", "".join(c for c in decomposed if not unicodedata.combining(c)))


def district_forms(district: str) -> tuple[str, str]:
    """(running text, upper case) genitive of a district name, e.g. («Λευκωσίας», «ΛΕΥΚΩΣΙΑΣ»)."""
    upper = _strip_accents(district.strip()).upper()
    return DISTRICT_GENITIVE.get(upper, upper.capitalize()), upper


def format_notice_date(value) -> str:
    return f"{value.day} {MONTHS_GENITIVE[value.month - 1]}, {value.year}"


def notice_filename(cycle) -> str:
    district = _district_of(cycle)
    return f"Γνωστοποίηση_{district_forms(district)[0]}_{cycle.announcement_date:%Y-%m-%d}.docx"


def _district_of(cycle) -> str:
    plots = list(cycle.plots.select_related("community").all())
    districts = sorted({plot.community.district for plot in plots})
    if not districts:
        fallback = getattr(cycle, "_fallback_community", None)
        if fallback is not None:
            return fallback.district
        raise NoticeExportError("Η γνωστοποίηση δεν προσφέρει κανένα οικόπεδο (9.1).")
    if len(districts) > 1:
        raise NoticeExportError(
            "Η γνωστοποίηση εκδίδεται από τον Έπαρχο μίας Επαρχίας (Ενότ. 4)· "
            f"βρέθηκαν: {', '.join(districts)}."
        )
    return districts[0]


def _plots_per_community(cycle) -> list[tuple[str, int]]:
    """(community, number of offered plots), by name; derived from offered plots."""
    plots = list(cycle.plots.select_related("community").all())
    counts = Counter(plot.community_id for plot in plots)
    communities = {plot.community_id: plot.community for plot in plots}
    rows = [
        (community.name, counts[community_id])
        for community_id, community in communities.items()
        if counts[community_id]
    ]
    rows.sort(key=lambda row: _strip_accents(row[0]).casefold())
    return rows


def _replace_everywhere(document, old: str, new: str) -> int:
    return sum(replace_in_paragraph(paragraph, old, new) for paragraph in document.paragraphs)


def _set_cell(cell, text: str) -> None:
    runs = cell.paragraphs[0].runs
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


def _clone_row(prototype):
    """A copy of a table row; Word wants every paragraph id (w14:paraId) to be unique."""
    tr = copy.deepcopy(prototype)
    for element in tr.iter():
        for attribute in (W14 + "paraId", W14 + "textId"):
            element.attrib.pop(attribute, None)
    return tr


def _fill_table(document, rows: list[tuple[str, int]]) -> None:
    """Keep the header row, then one row per community (Α/Α, Κοινότητα, number of plots)."""
    table = document.tables[0]
    prototype = copy.deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for number, (community, plots) in enumerate(rows, start=1):
        table._tbl.append(_clone_row(prototype))
        cells = table.rows[-1].cells
        _set_cell(cells[0], str(number))
        _set_cell(cells[1], community)
        _set_cell(cells[2], str(plots))


def render_notice(cycle) -> BytesIO:
    """The announcement as a Πρότυπο 1 notice, ready to be sent as a download."""
    district = _district_of(cycle)
    rows = _plots_per_community(cycle)
    if not rows:
        raise NoticeExportError("Η γνωστοποίηση δεν προσφέρει κανένα οικόπεδο (9.1).")
    period = PERIOD_IN_MONTHS.get(cycle.submission_months)
    if period is None:
        raise NoticeExportError("Η προθεσμία υποβολής πρέπει να είναι από 1 έως 12 μήνες.")
    genitive, genitive_upper = district_forms(district)

    document = Document(TEMPLATE_PATH)
    replacements = (
        (SAMPLE_DATE, format_notice_date(cycle.announcement_date)),
        (SAMPLE_DISTRICT_UPPER, genitive_upper),
        (SAMPLE_DISTRICT, genitive),
        (SAMPLE_PERIOD, period),
        (SAMPLE_END_DATE, f"{cycle.submission_end_date:%d/%m/%Y}"),
    )
    for old, new in replacements:
        if not _replace_everywhere(document, old, new):
            raise NoticeExportError(f"Το πρότυπο της γνωστοποίησης δεν περιέχει το «{old}».")
    _fill_table(document, rows)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
