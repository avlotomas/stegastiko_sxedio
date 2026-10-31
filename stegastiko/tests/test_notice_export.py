"""9.3 «Export»: the Πρότυπο 1 notice filled in from an announcement."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from docx import Document

from core.models import Community
from plots.models import Plot
from plots.notice_export import (
    NoticeExportError,
    district_forms,
    format_notice_date,
    render_notice,
    replace_in_paragraph,
)
from tests.fixtures.cycles import make_cycle

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _community(name, district="ΛΕΥΚΩΣΙΑΣ", code="1"):
    return Community.objects.create(
        district=district, name=name, contact_email=f"{code}@example.org", community_folder_code=code
    )


def _plots(community, count):
    return [
        Plot.objects.create(
            community=community,
            design_number=f"{community.name}-{n}",
            valuation_amount=Decimal("50000"),
            final_area_sqm=Decimal("500"),
        )
        for n in range(count)
    ]


def _cycle(*offers, **fields):
    """An announcement of the communities of `offers`, each given as (community, plot count)."""
    values = {
        "announcement_date": date(2026, 10, 5),
        "submission_months": 3,
    }
    values.update(fields)
    cycle = make_cycle(*[community for community, _ in offers], **values)
    for community, count in offers:
        cycle.plots.add(*_plots(community, count))
    return cycle


def _text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for row in document.tables[0].rows:
        parts.append("|".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _rows(document):
    return [[cell.text for cell in row.cells] for row in document.tables[0].rows]


@pytest.fixture
def staff_client(client, make_user):
    client.force_login(make_user("supervisor", "Προϊστάμενος ελέγχου"))
    return client


@pytest.mark.django_db
def test_notice_carries_the_announcement_data():
    abg = _community("Άγιος Γεώργιος", code="2")
    cycle = _cycle((abg, 5))

    document = Document(render_notice(cycle))
    text = _text(document)

    assert "5 Οκτωβρίου, 2026" in text
    assert "τριών (3) μηνών" in text and "05/01/2027" in text
    assert [row[1:] for row in _rows(document)[1:]] == [["Άγιος Γεώργιος", "5"]]
    assert [row[0] for row in _rows(document)[1:]] == ["1"]
    # Nothing of the sample notice is left.
    for sample in ("28 Σεπτεμβρίου", "30/11/2026", "δύο (2)", "Αγροκηπιά", "Μένοικο"):
        assert sample not in text


@pytest.mark.django_db
def test_notice_uses_the_district_of_the_communities():
    limassol = _community("Άγιος Αθανάσιος", district="ΛΕΜΕΣΟΥ")
    document = Document(render_notice(_cycle((limassol, 2))))
    text = _text(document)
    assert "ΕΠΑΡΧΙΑΚΗΣ ΔΙΟΙΚΗΣΗΣ ΛΕΜΕΣΟΥ" in text
    assert "Ο Έπαρχος Λεμεσού" in text
    assert "Επαρχίας Λεμεσού" in text
    assert "ΛΕΥΚΩΣΙΑΣ" not in text
    assert "Επαρχίας Λευκωσίας" not in text


@pytest.mark.django_db
def test_a_community_without_offered_plots_is_not_listed():
    listed = _community("Α", code="1")
    empty = _community("Β", code="2")
    cycle = _cycle((listed, 1), (empty, 0))
    assert [row[1] for row in _rows(Document(render_notice(cycle)))[1:]] == ["Α"]


@pytest.mark.django_db
def test_notice_without_plots_is_refused():
    cycle = _cycle((_community("Α"), 0))
    with pytest.raises(NoticeExportError):
        render_notice(cycle)


@pytest.mark.django_db
def test_export_view_downloads_a_docx(staff_client):
    cycle = _cycle((_community("Α"), 2), submission_months=1)
    response = staff_client.get(reverse("announcements:export", args=[cycle.pk]))
    assert response.status_code == 200
    assert response["Content-Type"] == DOCX
    assert "attachment" in response["Content-Disposition"]
    document = Document(BytesIO(b"".join(response.streaming_content)))
    assert "ενός (1) μηνός" in _text(document)


@pytest.mark.django_db
def test_export_view_without_plots_goes_back_with_a_message(staff_client):
    cycle = _cycle((_community("Α"), 0))
    response = staff_client.get(reverse("announcements:export", args=[cycle.pk]), follow=True)
    assert response.redirect_chain[-1][0] == reverse("announcements:detail", args=[cycle.pk])
    assert "κανένα οικόπεδο" in response.content.decode()


@pytest.mark.django_db
def test_export_button_is_on_the_detail_page(staff_client):
    cycle = _cycle((_community("Α"), 1))
    html = staff_client.get(reverse("announcements:detail", args=[cycle.pk])).content.decode()
    assert reverse("announcements:export", args=[cycle.pk]) in html
    assert "data-export-announcement" in html


def test_export_needs_login(client):
    response = client.get(reverse("announcements:export", args=[1]))
    assert response.status_code == 302 and "login" in response["Location"]


def test_replace_in_paragraph_handles_text_split_over_runs():
    paragraph = Document().add_paragraph()
    for text in ("Ο Έπαρχος ", "Λευ", "κωσίας", " ανακοινώνει"):
        paragraph.add_run(text)
    assert replace_in_paragraph(paragraph, "Λευκωσίας", "Λεμεσού") == 1
    assert paragraph.text == "Ο Έπαρχος Λεμεσού ανακοινώνει"


def test_replace_in_paragraph_does_not_loop_when_the_new_text_contains_the_old():
    paragraph = Document().add_paragraph("ab")
    assert replace_in_paragraph(paragraph, "a", "aa") == 1
    assert paragraph.text == "aab"


@pytest.mark.django_db
def test_end_date_is_computed_from_the_months():
    cycle = make_cycle(_community("Α"), announcement_date=date(2026, 9, 28), submission_months=2)
    assert cycle.submission_end_date == date(2026, 11, 28)
    cycle.announcement_date, cycle.submission_months = date(2026, 12, 31), 2
    cycle.save()
    assert cycle.submission_end_date == date(2027, 2, 28)
    cycle.announcement_date, cycle.submission_months = date(2026, 1, 31), 1
    cycle.save(update_fields=["announcement_date", "submission_months"])
    cycle.refresh_from_db()
    assert cycle.submission_end_date == date(2026, 2, 28)


def test_district_forms_and_dates():
    assert district_forms("ΛΕΥΚΩΣΙΑΣ") == ("Λευκωσίας", "ΛΕΥΚΩΣΙΑΣ")
    assert district_forms("Λάρνακας") == ("Λάρνακας", "ΛΑΡΝΑΚΑΣ")
    # An unknown district cannot get its accents back, but is still spelled in running text.
    assert district_forms("ΝΕΑΣ") == ("Νεας", "ΝΕΑΣ")
    assert format_notice_date(date(2026, 9, 28)) == "28 Σεπτεμβρίου, 2026"
