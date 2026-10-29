"""Shared Word (.docx) merge and placeholder replacement utilities.

Used by:
- 9.3 Notice export (`plots.notice_export`)
- Stage 5 Agreement template rendering (Παράρτημα 4, Ενότ. 12 — `allocation.AgreementTemplate`)
"""

from __future__ import annotations

import copy
import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Iterable

from django.core.exceptions import ValidationError
from django.utils import timezone
from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

from core.services import get_setting

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

APPENDIX4_TEMPLATE_PATH = Path(__file__).with_name("docx_templates") / "appendix4_template.docx"

FILLED_FIELD_FONT_NAME = "Times New Roman"

AGREEMENT_DISTRICT_OFFICER_KEY = "agreementDistrictOfficer"
AGREEMENT_DISTRICT_OFFICER_DEFAULT = "Λευκωσίας"

DISTRICT_GENITIVE = {
    "ΛΕΥΚΩΣΙΑΣ": "Λευκωσίας",
    "ΛΕΜΕΣΟΥ": "Λεμεσού",
    "ΛΑΡΝΑΚΑΣ": "Λάρνακας",
    "ΠΑΦΟΥ": "Πάφου",
    "ΑΜΜΟΧΩΣΤΟΥ": "Αμμοχώστου",
    "ΚΕΡΥΝΕΙΑΣ": "Κερύνειας",
}

AGREEMENT_PLACEHOLDERS: tuple[tuple[str, str], ...] = (
    ("folder_number", "Αρ. Φακέλου (Αριθμός πρωτοκόλλου αίτησης)"),
    ("protocol_number", "Αριθμός πρωτοκόλλου αίτησης"),
    ("submitted_on", "Ημερομηνία υποβολής αίτησης"),
    ("announcement_date", "Ημερομηνία γνωστοποίησης"),
    ("community", "Κοινότητα / Δημοτικό Διαμέρισμα αίτησης"),
    ("district", "Επαρχία"),
    ("district_officer", "Έπαρχος (παραμετρικό πεδίο στις Ρυθμίσεις)"),
    ("location", "Τόπος υπογραφής Συμφωνίας (αρχικά: Επαρχιακή Διοίκηση)"),
    ("applicant_name", "Ονοματεπώνυμο αιτητή/αιτήτριας"),
    ("applicant_first_name", "Όνομα αιτητή/αιτήτριας"),
    ("applicant_last_name", "Επίθετο αιτητή/αιτήτριας"),
    ("applicant_identity_number", "Αρ. Δελτίου Ταυτότητας αιτητή/αιτήτριας"),
    ("spouse_name", "Ονοματεπώνυμο συζύγου/συμβίου (αν υπάρχει)"),
    ("spouse_first_name", "Όνομα συζύγου/συμβίου"),
    ("spouse_last_name", "Επίθετο συζύγου/συμβίου"),
    ("spouse_identity_number", "Αρ. Δελτίου Ταυτότητας συζύγου/συμβίου"),
    ("parties_names", "Ονοματεπώνυμα συμβαλλομένων δικαιούχων (Ζεύγος ή Μονογονιός, Ενότ. 12)"),
    ("family_status", "Οικογενειακή κατάσταση"),
    ("address", "Διεύθυνση διαμονής αιτητή (από την αίτηση παραχώρησης οικοπέδου)"),
    ("email", "Ηλεκτρονική διεύθυνση"),
    ("phone", "Τηλέφωνο επικοινωνίας"),
    ("plot_number", "Αριθμός οικοπέδου (Σχεδιασμού / ΤΚΧ / Τίτλου)"),
    ("parcel_number", "Αριθμός τεμαχίου"),
    ("sheet_plan", "Φύλλο / Σχέδιο"),
    ("plot_area_sqm", "Εμβαδόν οικοπέδου (τ.μ.)"),
    ("valuation_amount", "Αξία Γενικής Εκτίμησης (€, Ενότ. 13)"),
    ("disposal_price", "Τιμή διάθεσης (25% Γενικής Εκτίμησης, €, Ενότ. 13, 15)"),
    ("price", "Τιμή διάθεσης (€)"),
    ("agreement_date", "Ημερομηνία υπογραφής Συμφωνίας"),
    ("appointment_date", "Ημερομηνία προσέλευσης"),
    ("appointment_datetime", "Ημερομηνία/ώρα προσέλευσης"),
    ("agreement_deadline", "Προθεσμία υπογραφής/πληρωμής"),
    ("rank", "Σειρά προτεραιότητας δικαιούχου"),
    ("lottery_date", "Ημερομηνία κλήρωσης οικοπέδων"),
)

ALLOWED_AGREEMENT_PLACEHOLDERS: frozenset[str] = frozenset(
    key for key, _ in AGREEMENT_PLACEHOLDERS
)


def style_filled_run(run: Run, font_name: str = FILLED_FIELD_FONT_NAME) -> None:
    """Style a filled-in run with bold, italic, and a distinct font family."""
    run.bold = True
    run.italic = True
    run.font.name = font_name
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font_name)


def replace_segments_in_paragraph(
    paragraph, old: str, segments: Iterable[tuple[str, bool]]
) -> int:
    """Replace `old` in `paragraph` with `segments` of `(text, is_filled)`.

    Segments with `is_filled=True` are placed in their own run and styled with
    bold, italic, and `FILLED_FIELD_FONT_NAME` so filled values stand out from
    the surrounding template text.
    """
    seg_list = [(str(text), bool(is_filled)) for text, is_filled in segments if text]
    replacement_len = sum(len(text) for text, _ in seg_list)
    replaced = 0
    search_from = 0
    while True:
        runs = list(paragraph.runs)
        if not runs:
            return replaced
        texts = [run.text for run in runs]
        full_text = "".join(texts)
        start = full_text.find(old, search_from)
        if start < 0:
            return replaced
        end = start + len(old)

        overlapping: list[tuple[Run, str, int, int]] = []
        position = 0
        for run, text in zip(runs, texts):
            run_start, position = position, position + len(text)
            if position <= start or run_start >= end:
                continue
            overlapping.append((run, text, run_start, position))

        if not overlapping:
            return replaced

        r_first, text_first, start_first, _ = overlapping[0]
        r_last, text_last, start_last, _ = overlapping[-1]
        prefix_text = text_first[: start - start_first]
        suffix_text = text_last[end - start_last :]

        to_insert: list[tuple[str, bool, object]] = []
        if prefix_text:
            to_insert.append((prefix_text, False, r_first._r))
        for seg_text, is_filled in seg_list:
            to_insert.append((seg_text, is_filled, r_first._r))
        if suffix_text:
            to_insert.append((suffix_text, False, r_last._r))

        if not to_insert:
            to_insert.append(("", False, r_first._r))

        for seg_text, is_filled, template_r in to_insert:
            new_r_el = copy.deepcopy(template_r)
            r_first._r.addprevious(new_r_el)
            new_run = Run(new_r_el, paragraph)
            new_run.text = seg_text
            if is_filled:
                style_filled_run(new_run)

        parent = r_first._r.getparent()
        for run, _, _, _ in overlapping:
            parent.remove(run._r)

        replaced += 1
        search_from = start + replacement_len


def replace_in_paragraph(
    paragraph, old: str, new: str, *, highlight: bool = False
) -> int:
    """Replace `old` by `new` in a paragraph even where Word split it over several runs.

    When `highlight=True`, the replacement run is styled with bold, italic, and
    `FILLED_FIELD_FONT_NAME`. Returns the number of replacements.
    """
    if highlight:
        return replace_segments_in_paragraph(paragraph, old, [(new, True)])
    replaced = 0
    search_from = 0
    while True:
        runs = paragraph.runs
        texts = [run.text for run in runs]
        start = "".join(texts).find(old, search_from)
        if start < 0:
            return replaced
        end = start + len(old)
        position = 0
        first = True
        for run, text in zip(runs, texts):
            run_start, position = position, position + len(text)
            if position <= start or run_start >= end:
                continue
            low = max(start, run_start) - run_start
            high = min(end, position) - run_start
            run.text = text[:low] + (new if first else "") + text[high:]
            first = False
        replaced += 1
        search_from = start + len(new)


def iter_document_paragraphs(document: Document) -> Iterable:
    """Yield every paragraph in body, tables, headers, and footers of a Word document."""
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for container in (section.header, section.footer):
            if container is None:
                continue
            yield from container.paragraphs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def replace_in_document(document: Document, replacements: dict[str, str]) -> int:
    total = 0
    for paragraph in iter_document_paragraphs(document):
        for old, new in replacements.items():
            total += replace_in_paragraph(paragraph, old, new)
    return total


def extract_placeholders(file_or_bytes) -> set[str]:
    """Return all `{{placeholder}}` names found in a `.docx` file."""
    if isinstance(file_or_bytes, (bytes, bytearray)):
        stream = BytesIO(file_or_bytes)
    else:
        pos = file_or_bytes.tell() if hasattr(file_or_bytes, "tell") else None
        if hasattr(file_or_bytes, "seek"):
            file_or_bytes.seek(0)
        data = file_or_bytes.read()
        if pos is not None and hasattr(file_or_bytes, "seek"):
            file_or_bytes.seek(pos)
        stream = BytesIO(data)
    try:
        doc = Document(stream)
    except Exception as exc:
        raise ValidationError("Το αρχείο δεν είναι έγκυρο έγγραφο Word (.docx).") from exc
    found: set[str] = set()
    for paragraph in iter_document_paragraphs(doc):
        text = "".join(run.text for run in paragraph.runs) or paragraph.text
        for match in PLACEHOLDER_PATTERN.finditer(text):
            found.add(match.group(1))
    return found


def validate_agreement_template_file(file_or_bytes) -> set[str]:
    """Validate that a `.docx` file is readable and only uses supported placeholders."""
    found = extract_placeholders(file_or_bytes)
    unknown = sorted(found - ALLOWED_AGREEMENT_PLACEHOLDERS)
    if unknown:
        formatted = ", ".join(f"{{{{{name}}}}}" for name in unknown)
        raise ValidationError(
            f"Το πρότυπο περιέχει μη υποστηριζόμενα πεδία: {formatted}."
        )
    return found


def _district_genitive(district: str) -> str:
    if not district:
        return ""
    decomposed = unicodedata.normalize("NFD", district.strip())
    upper = unicodedata.normalize(
        "NFC", "".join(c for c in decomposed if not unicodedata.combining(c))
    ).upper()
    return DISTRICT_GENITIVE.get(upper, upper.capitalize())


def _date_str(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "tzinfo") and value.tzinfo is not None:
        value = timezone.localtime(value)
    return f"{value:%d/%m/%Y}"


def build_agreement_context(allocation) -> dict[str, str]:
    """Build the placeholder dictionary for an `Allocation` (Παράρτημα 4, Ενότ. 12, 13, 15)."""
    app = allocation.application
    plot = allocation.parcel
    cycle = (
        app.announcement
        if (app is not None and getattr(app, "announcement_id", None))
        else (
            allocation.lottery.submission_cycle
            if getattr(allocation, "lottery_id", None) and allocation.lottery is not None
            else None
        )
    )
    community = (
        (app.community if app is not None and app.community is not None else None)
        or (plot.community if plot is not None else None)
        or (cycle.community if cycle is not None else None)
    )
    applicant = app.applicant if (app is not None and getattr(app, "applicant_id", None)) else None
    spouse = (
        app.spouse
        if (app is not None and app.is_couple and getattr(app, "spouse_id", None))
        else None
    )
    applicant_party = app.applicant_party if app is not None else None

    applicant_name = f"{applicant.first_name} {applicant.last_name}".strip() if applicant else ""
    spouse_name = f"{spouse.first_name} {spouse.last_name}".strip() if spouse else ""
    if spouse_name:
        parties_names = f"{applicant_name} και {spouse_name}"
    else:
        parties_names = applicant_name

    try:
        agreement = allocation.agreement
    except Exception:
        agreement = None
    try:
        payment = allocation.payment
    except Exception:
        payment = None

    price_val = (
        payment.amount
        if payment is not None and payment.amount is not None
        else (plot.disposal_price if plot is not None else None)
    )
    price_str = f"{price_val:.2f}" if price_val is not None else "0.00"
    valuation_str = (
        f"{plot.valuation_amount:.2f}"
        if plot is not None and plot.valuation_amount is not None
        else "0.00"
    )
    area_str = (
        f"{plot.final_area_sqm:.2f}"
        if plot is not None and plot.final_area_sqm is not None
        else ""
    )
    signed_date = (
        agreement.signed_on
        if agreement is not None and agreement.signed_on is not None
        else timezone.localdate()
    )

    appt_dt = None
    if agreement is not None and agreement.appointment_at is not None:
        appt_dt = agreement.appointment_at
    elif cycle is not None and cycle.agreement_appointment_at is not None:
        appt_dt = cycle.agreement_appointment_at

    if appt_dt is not None:
        dt = timezone.localtime(appt_dt) if timezone.is_aware(appt_dt) else appt_dt
        appointment_date_str = f"{dt:%d/%m/%Y}"
        appointment_datetime_str = f"{dt:%d/%m/%Y %H:%M}"
    else:
        appointment_date_str = f"{signed_date:%d/%m/%Y}"
        appointment_datetime_str = appointment_date_str

    deadline_val = (
        agreement.deadline
        if agreement is not None and agreement.deadline is not None
        else (cycle.agreement_deadline if cycle is not None else None)
    )
    deadline_str = f"{deadline_val:%d/%m/%Y}" if deadline_val is not None else ""

    location_str = (
        (cycle.agreement_location.strip() if cycle is not None and cycle.agreement_location else "")
        or "Επαρχιακή Διοίκηση"
    )
    district_officer = (
        get_setting(AGREEMENT_DISTRICT_OFFICER_KEY, AGREEMENT_DISTRICT_OFFICER_DEFAULT) or ""
    ).strip()

    street_part = f"{app.residence_street} {app.residence_number}".strip() if app else ""
    if app and getattr(app, "residence_apartment", ""):
        street_part = f"{street_part}, {app.residence_apartment}".strip(", ")
    address_parts = [
        street_part,
        app.residence_postal_code if app else "",
        app.residence_community if app else "",
    ]
    address = ", ".join(p for p in address_parts if p)

    protocol_number = (app.protocol_number or "") if app else ""
    announcement_date_str = (
        f"{cycle.announcement_date:%d/%m/%Y}"
        if cycle is not None and cycle.announcement_date is not None
        else ""
    )
    submitted_on_str = (
        f"{app.submitted_on:%d/%m/%Y}" if app is not None and app.submitted_on is not None else ""
    )
    district_raw = community.district if community is not None else ""
    district_genitive = _district_genitive(district_raw) or district_officer

    # Approval letter date & priority rank (Προοίμιο Ε.2)
    rank_str = ""
    approval_letter_date_str = ""
    lottery_letter_date_str = ""
    if app is not None and getattr(app, "pk", None):
        entries = list(app.ranking_entries.select_related("ranking_list").all())
        final_entry = next(
            (e for e in entries if e.ranking_list.list_type == "final"), None
        )
        prelim_entry = next(
            (e for e in entries if e.ranking_list.list_type == "preliminary"), None
        )
        chosen_entry = final_entry or prelim_entry
        if chosen_entry is not None:
            rank_val = chosen_entry.finalized_rank or chosen_entry.proposed_rank
            if rank_val is not None:
                rank_str = str(rank_val)
        if not rank_str and app.priority_rank is not None:
            rank_str = str(app.priority_rank)

        ref_entry = prelim_entry or final_entry
        if ref_entry is not None:
            approval_letter_date_str = (
                _date_str(ref_entry.notified_at)
                or _date_str(ref_entry.ranking_list.notifications_sent_at)
                or _date_str(ref_entry.ranking_list.finalized_at)
            )
        if not approval_letter_date_str and cycle is not None:
            approval_letter_date_str = _date_str(cycle.applications_committee_date)

        if final_entry is not None:
            lottery_letter_date_str = (
                _date_str(final_entry.ranking_list.notifications_sent_at)
                or _date_str(final_entry.ranking_list.finalized_at)
            )

    lottery = getattr(allocation, "lottery", None)
    if lottery is None and cycle is not None and getattr(cycle, "pk", None):
        lottery = cycle.lotteries.order_by("-id").first()
    lottery_date_str = _date_str(lottery.lottery_date) if lottery is not None else ""
    if not lottery_letter_date_str:
        lottery_letter_date_str = lottery_date_str

    owner_email = (
        get_setting("smtpFromEmail", "noreply@moi.gov.cy")
        or (community.contact_email if community is not None else "")
        or "noreply@moi.gov.cy"
    )

    return {
        "folder_number": protocol_number,
        "protocol_number": protocol_number,
        "submitted_on": submitted_on_str,
        "announcement_date": announcement_date_str,
        "community": community.name if community else "",
        "district": district_raw,
        "district_genitive": district_genitive,
        "district_officer": district_officer,
        "location": location_str,
        "applicant_name": applicant_name,
        "applicant_first_name": applicant.first_name if applicant else "",
        "applicant_last_name": applicant.last_name if applicant else "",
        "applicant_identity_number": applicant.identity_number if applicant else "",
        "spouse_name": spouse_name,
        "spouse_first_name": spouse.first_name if spouse else "",
        "spouse_last_name": spouse.last_name if spouse else "",
        "spouse_identity_number": spouse.identity_number if spouse else "",
        "parties_names": parties_names,
        "family_status": app.get_family_status_display() if app else "",
        "address": address,
        "email": app.email if app else "",
        "owner_email": owner_email,
        "phone": applicant_party.mobile_phone if applicant_party else "",
        "plot_number": (plot.number or plot.row_label) if plot else "",
        "parcel_number": plot.parcel_number if plot else "",
        "sheet_plan": plot.sheet_plan if plot else "",
        "plot_area_sqm": area_str,
        "valuation_amount": valuation_str,
        "disposal_price": price_str,
        "price": price_str,
        "agreement_date": f"{signed_date:%d/%m/%Y}",
        "appointment_date": appointment_date_str,
        "appointment_datetime": appointment_datetime_str,
        "agreement_deadline": deadline_str,
        "rank": rank_str,
        "approval_letter_date": approval_letter_date_str,
        "lottery_letter_date": lottery_letter_date_str,
        "lottery_date": lottery_date_str,
    }


def _single_field_rule(
    prefix: str, dots: str, value: str, suffix: str = "", *, new_suffix: str | None = None
) -> tuple[str, list[tuple[str, bool]]] | None:
    if not value:
        return None
    out_suffix = suffix if new_suffix is None else new_suffix
    return (
        f"{prefix}{dots}{suffix}",
        [(prefix, False), (value, True), (out_suffix, False)],
    )


def _fill_appendix4_document(doc: Document, context: dict[str, str]) -> None:
    """Replace the dotted fields of the official Παράρτημα 4 (`appendix4.doc`) template."""
    if context["spouse_name"]:
        parties_segments: list[tuple[str, bool]] = [
            ("των ", False),
            (
                context["applicant_name"] or "……………………………………..",
                bool(context["applicant_name"]),
            ),
            (" με Α.Δ.Τ. ", False),
            (
                context["applicant_identity_number"] or "………………………….",
                bool(context["applicant_identity_number"]),
            ),
            (" και ", False),
            (context["spouse_name"], True),
            (" με Α.Δ.Τ. ", False),
            (
                context["spouse_identity_number"] or "………………………..",
                bool(context["spouse_identity_number"]),
            ),
        ]
        article_11_rule: tuple[str, list[tuple[str, bool]]] = (
            "Οι ……………………… και …………………………",
            [
                ("Οι ", False),
                (
                    context["applicant_name"] or "………………………",
                    bool(context["applicant_name"]),
                ),
                (" και ", False),
                (context["spouse_name"], True),
            ],
        )
    else:
        parties_segments = [
            ("των ", False),
            (
                context["applicant_name"] or "……………………………………..",
                bool(context["applicant_name"]),
            ),
            (" με Α.Δ.Τ. ", False),
            (
                context["applicant_identity_number"] or "………………………….",
                bool(context["applicant_identity_number"]),
            ),
        ]
        article_11_rule = (
            "Οι ……………………… και ………………………… (ονοματεπώνυμο Αιτητή/Αιτήτριας και Συζύγου/Συμβίου, ό,τι ισχύει), οι οποίοι έχει οριστεί να καλούνται, στο πλαίσιο της παρούσας Συμφωνίας, ως «ο Δικαιούχος», υπέχουν",
            [
                ("Ο/Η ", False),
                (
                    context["applicant_name"] or "………………………",
                    bool(context["applicant_name"]),
                ),
                (
                    " (ονοματεπώνυμο Αιτητή/Αιτήτριας), ο/η οποίος/α έχει οριστεί να καλείται, "
                    "στο πλαίσιο της παρούσας Συμφωνίας, ως «ο Δικαιούχος», υπέχει",
                    False,
                ),
            ],
        )

    raw_rules: list[tuple[str, list[tuple[str, bool]]] | None] = [
        # Αρ. Φακέλου
        _single_field_rule("Αρ. Φακ. ", "......................", context["folder_number"]),
        # Συμφωνία που γίνεται σήμερα στις : ημερομηνία προσέλευσης / στην : τόπος
        _single_field_rule(
            "Συμφωνία που γίνεται σήμερα στις ",
            ".....................",
            context["appointment_date"],
        ),
        _single_field_rule("στην ", "……………..", context["location"]),
        # του Επάρχου : παραμετρικό πεδίο από τα settings
        _single_field_rule(
            "του Επάρχου ",
            ".................................",
            context["district_officer"],
        ),
        # και των : όνομα/επίθετο αιτητή (και συζύγου) / Α.Δ.Τ. / διεύθυνση διαμονής
        (
            "των ……………………………………..με Α.Δ.Τ. ………………………….και ……………………………….. με Α.Δ.Τ. ………………………..",
            parties_segments,
        ),
        _single_field_rule(
            "με μόνιμη διεύθυνση διαμονής στην οδό ",
            "…………………………………",
            context["address"],
        ),
        # Γ. ΕΠΕΙΔΗ ο/η Έπαρχος / ανακοίνωσε στις : ημερομηνία γνωστοποίησης
        _single_field_rule("ΕΠΕΙΔΗ ο/η Έπαρχος ", "………….", context["district_officer"]),
        _single_field_rule(
            "ανακοίνωσε στις ",
            "………/…………./………..",
            context["announcement_date"],
            new_suffix=" ",
        ),
        # Δ. ο Δικαιούχος υπέβαλε στις / έλαβε τον αριθμό αίτησης
        _single_field_rule(
            "ο Δικαιούχος υπέβαλε στις ",
            "………../…………/……………",
            context["submitted_on"],
        ),
        _single_field_rule(
            "έλαβε τον αριθμό αίτησης ",
            "…………….",
            context["protocol_number"],
        ),
        # Ε.1 Κοινότητα/Δημοτικό Διαμέρισμα : κοινότητα για την οποία έγινε η αίτηση
        _single_field_rule(
            "Κοινότητα/Δημοτικό Διαμέρισμα ",
            "………………..",
            context["community"],
            ",",
        ),
        # Ε.2 Έγκριση αίτησης & σειρά προτεραιότητας
        _single_field_rule(
            "στις ",
            "………./…………./………..",
            context["approval_letter_date"],
            " ο Δικαιούχος",
        ),
        _single_field_rule("Έπαρχο ", "……………", context["district_officer"]),
        _single_field_rule(
            "με επιστολή ημερομηνίας ",
            "…./……./…….",
            context["approval_letter_date"],
        ),
        _single_field_rule("υπ’ αριθμό φακέλου ", "………….", context["folder_number"]),
        _single_field_rule(
            "σειρά προτεραιότητας υπ’ αριθμόν ",
            "………….",
            context["rank"],
            ",",
        ),
        _single_field_rule(
            "επιστολή ημερομηνίας",
            "…./……./…….",
            f" {context['approval_letter_date']}" if context["approval_letter_date"] else "",
        )
        if False
        else (
            (
                "επιστολή ημερομηνίας…./……./…….",
                [("επιστολή ημερομηνίας ", False), (context["approval_letter_date"], True)],
            )
            if context["approval_letter_date"]
            else None
        ),
        # ΣΤ. Πρόσκληση σε κλήρωση
        _single_field_rule(
            "ο Έπαρχος ",
            "……………",
            context["district_officer"],
            "ενημέρωσε",
            new_suffix=" ενημέρωσε",
        ),
        _single_field_rule(
            "επιστολή του ημερομηνίας ",
            "……./………./………..",
            context["lottery_letter_date"],
        ),
        _single_field_rule(
            "υπ’ αριθμόν φακέλου ",
            "………..",
            context["folder_number"],
            "(",
            new_suffix=" (",
        ),
        _single_field_rule(
            "επιστολή ημερομηνίας ",
            "…./……./…….",
            context["lottery_letter_date"],
            "επισυνάπτεται",
            new_suffix=" επισυνάπτεται",
        ),
        _single_field_rule(
            "ότι στις ",
            "………/……/………",
            context["lottery_date"],
            " θα διεξαγόταν",
        ),
        _single_field_rule(
            "Επαρχιακής Διοίκησης ",
            "………………..",
            context["district_genitive"],
            "στα πλαίσια",
            new_suffix=" στα πλαίσια",
        ),
        # Ζ. Κλήρωση & στοιχεία οικοπέδου
        _single_field_rule(
            "στις ",
            "…./……./…….",
            context["lottery_date"],
            " διεξήχθη",
        ),
        _single_field_rule(
            "Επαρχιακής Διοίκησης ",
            "…………",
            context["district_genitive"],
            " η ανωτέρω",
        ),
        _single_field_rule(
            "Αριθμόν Εγγραφής ",
            "………….",
            context["plot_number"],
            ",",
        ),
        _single_field_rule(
            "Φύλλο/ Σχέδιο ",
            "……/………..",
            context["sheet_plan"],
            ", ",
        ),
        _single_field_rule(
            "Τμήμα ",
            "…………….",
            context["parcel_number"],
            ",",
        ),
        _single_field_rule(
            "Περιοχή ",
            "…………",
            context["community"],
            ",",
        ),
        (
            "Κοινότητα/Δημοτικό Διαμέρισμα ………………… στην Επαρχία ………………,",
            [
                ("Κοινότητα/Δημοτικό Διαμέρισμα ", False),
                (context["community"] or "…………………", bool(context["community"])),
                (" στην Επαρχία ", False),
                (
                    context["district_genitive"] or "………………",
                    bool(context["district_genitive"]),
                ),
                (",", False),
            ],
        ),
        _single_field_rule(
            "έκτασης ",
            "……………………….",
            f"{context['plot_area_sqm']} τ.μ." if context["plot_area_sqm"] else "",
            "(στο εξής",
            new_suffix=" (στο εξής",
        ),
        # Άρθρο 2 & Άρθρο 3.1: Αξία Οικοπέδου & 25% Τιμή Διάθεσης
        (
            "€………………… (……………………………..…..Ευρώ)",
            [
                ("€", False),
                (context["valuation_amount"], True),
                (" (", False),
                (context["valuation_amount"], True),
                (" Ευρώ)", False),
            ],
        ),
        (
            "€…………………………….. (………………………………Ευρώ)",
            [
                ("€", False),
                (context["disposal_price"], True),
                (" (", False),
                (context["disposal_price"], True),
                (" Ευρώ)", False),
            ],
        ),
        # Άρθρο 11 & Κατακλείδα
        article_11_rule,
        _single_field_rule("", "……./…………./………..", context["appointment_date"]),
    ]

    rules: list[tuple[str, list[tuple[str, bool]]]] = [r for r in raw_rules if r is not None]

    contact_section = ""
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "(α)\tΣτον Ιδιοκτήτη" in text or "Στον Ιδιοκτήτη:" in text:
            contact_section = "owner"
        elif "(β)\tΣτον Δικαιούχο" in text or "Στον Δικαιούχο:" in text:
            contact_section = "beneficiary"
        elif text.startswith("9.2"):
            contact_section = ""

        if contact_section == "owner":
            if "Στη διεύθυνση: …………………………………." in text and context["location"]:
                replace_in_paragraph(paragraph, "………………………………….", context["location"], highlight=True)
                continue
            if "Στο Ηλεκτρονικό Ταχυδρομείο: ………………………" in text and context["owner_email"]:
                replace_in_paragraph(paragraph, "………………………", context["owner_email"], highlight=True)
                continue
        elif contact_section == "beneficiary":
            if "Στη διεύθυνση: …………………………………." in text and context["address"]:
                replace_in_paragraph(paragraph, "………………………………….", context["address"], highlight=True)
                continue
            if "Στο Ηλεκτρονικό Ταχυδρομείο: ………………………" in text and context["email"]:
                replace_in_paragraph(paragraph, "………………………", context["email"], highlight=True)
                continue

        for old, segs in rules:
            replace_segments_in_paragraph(paragraph, old, segs)

    # Signature table (Table 0: ΙΔΙΟΚΤΗΤΗΣ / ΔΙΚΑΙΟΥΧΟΣ)
    if doc.tables:
        sig_table = doc.tables[0]
        if len(sig_table.rows) >= 1 and len(sig_table.rows[0].cells) >= 2:
            owner_cell = sig_table.rows[0].cells[0]
            for p in owner_cell.paragraphs:
                if context["district_officer"]:
                    replace_segments_in_paragraph(
                        p,
                        "Επάρχου ………………….",
                        [("Επάρχου ", False), (context["district_officer"], True)],
                    )
            beneficiary_cell = sig_table.rows[0].cells[1]
            name_slots = [
                p for p in beneficiary_cell.paragraphs if p.text.strip() == "……………………."
            ]
            id_slots = [
                p for p in beneficiary_cell.paragraphs if "Α.Δ.Τ. ………………." in p.text
            ]
            if name_slots and context["applicant_name"]:
                replace_in_paragraph(
                    name_slots[0], "…………………….", context["applicant_name"], highlight=True
                )
            if id_slots and context["applicant_identity_number"]:
                replace_segments_in_paragraph(
                    id_slots[0],
                    "Α.Δ.Τ. ……………….",
                    [("Α.Δ.Τ. ", False), (context["applicant_identity_number"], True)],
                )
            if len(name_slots) >= 2 and context["spouse_name"]:
                replace_in_paragraph(
                    name_slots[1], "…………………….", context["spouse_name"], highlight=True
                )
            if len(id_slots) >= 2 and context["spouse_identity_number"]:
                replace_segments_in_paragraph(
                    id_slots[1],
                    "Α.Δ.Τ. ……………….",
                    [("Α.Δ.Τ. ", False), (context["spouse_identity_number"], True)],
                )


def _build_default_agreement_document() -> Document:
    """Load the official Παράρτημα 4 Word template (`appendix4.doc` -> `appendix4_template.docx`)."""
    return Document(APPENDIX4_TEMPLATE_PATH)


def render_agreement_docx(allocation, template=None) -> BytesIO:
    """Fill the active `AgreementTemplate` (or default Παράρτημα 4) with `allocation` data."""
    if template is None:
        from allocation.models import AgreementTemplate

        template = AgreementTemplate.objects.filter(is_active=True).order_by("-id").first()

    if template is not None and template.template_file:
        template.template_file.open("rb")
        try:
            data = template.template_file.read()
        finally:
            template.template_file.close()
        doc = Document(BytesIO(data))
    else:
        doc = _build_default_agreement_document()

    context = build_agreement_context(allocation)
    _fill_appendix4_document(doc, context)
    for paragraph in iter_document_paragraphs(doc):
        full_text = "".join(run.text for run in paragraph.runs)
        for match in list(PLACEHOLDER_PATTERN.finditer(full_text)):
            token = match.group(0)
            key = match.group(1)
            if key in context:
                replace_in_paragraph(paragraph, token, str(context[key]), highlight=True)

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
