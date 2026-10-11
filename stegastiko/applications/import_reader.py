"""Read operator Excel import workbook (§9.2)."""

import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from applications.import_template import TEMPLATE_VERSION

KEY_ROW = 2
DATA_ROW = 4
TABLE_DATA_START = 4

# Uploaded workbooks are a single application; anything larger is not a template.
MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024

INVALID_WORKBOOK_MESSAGE = (
    "Το αρχείο δεν είναι έγκυρο αρχείο Excel (.xlsx) ή είναι κατεστραμμένο. "
    "Χρησιμοποιήστε το πρότυπο εισαγωγής."
)


@dataclass
class ParsedImport:
    metadata: dict[str, Any]
    fields: dict[str, Any]
    children: list[dict[str, Any]] = field(default_factory=list)
    income_rows: list[dict[str, Any]] = field(default_factory=list)
    supporting_documents: list[dict[str, Any]] = field(default_factory=list)


def _cell_value(value: Any) -> str:
    """Normalise every cell to text (ID numbers typed as numbers keep their digits)."""
    if value is None:
        return ""
    # datetime is a subclass of date: check it first.
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "Ναι" if value else "Όχι"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)), "f")
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _row_values(ws, row_number: int) -> list[Any]:
    for row in ws.iter_rows(min_row=row_number, max_row=row_number, values_only=True):
        return list(row)
    return []


def _read_metadata(ws) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        key = row[0]
        if key is None or str(key).strip() == "":
            continue
        data[str(key).strip()] = _cell_value(row[2] if len(row) > 2 else None)
    return data


def _read_horizontal(ws) -> dict[str, Any]:
    keys = _row_values(ws, KEY_ROW)
    values = _row_values(ws, DATA_ROW)
    data: dict[str, Any] = {}
    for idx, key in enumerate(keys):
        if key is None or str(key).strip() == "":
            continue
        data[str(key).strip()] = _cell_value(values[idx] if idx < len(values) else None)
    return data


def _read_table(ws) -> list[dict[str, Any]]:
    keys = [
        (idx, str(k).strip())
        for idx, k in enumerate(_row_values(ws, KEY_ROW))
        if k is not None and str(k).strip() != ""
    ]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=TABLE_DATA_START, values_only=True):
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            break
        rows.append(
            {key: _cell_value(row[idx] if idx < len(row) else None) for idx, key in keys}
        )
    return rows


def _load_workbook(source):
    size = getattr(source, "size", None)
    if size is not None and size > MAX_IMPORT_FILE_SIZE:
        raise ValidationError(
            f"Το αρχείο υπερβαίνει το μέγιστο επιτρεπτό μέγεθος "
            f"({MAX_IMPORT_FILE_SIZE // (1024 * 1024)} MB)."
        )
    try:
        return load_workbook(source, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError, TypeError):
        raise ValidationError(INVALID_WORKBOOK_MESSAGE)


def parse_import_workbook(source) -> ParsedImport:
    wb = _load_workbook(source)
    try:
        return _parse_loaded_workbook(wb)
    except ValidationError:
        raise
    except Exception:  # malformed sheet content (e.g. broken XML inside a valid zip)
        raise ValidationError(INVALID_WORKBOOK_MESSAGE)
    finally:
        wb.close()


def _parse_loaded_workbook(wb) -> ParsedImport:
    required = ("Μεταδεδομένα", "Αίτηση", "Πρόσωπο1", "Διαμονή", "Δηλώσεις")
    for name in required:
        if name not in wb.sheetnames:
            raise ValidationError(f"Λείπει φύλλο «{name}».")

    metadata = _read_metadata(wb["Μεταδεδομένα"])
    version = str(metadata.get("template_version", "")).strip()
    if version != TEMPLATE_VERSION:
        raise ValidationError(
            f"Άγνωστη έκδοση προτύπου: «{version}» (αναμενόμενη {TEMPLATE_VERSION})."
        )

    fields: dict[str, Any] = {}
    fields.update(_read_horizontal(wb["Αίτηση"]))
    fields.update(_read_horizontal(wb["Πρόσωπο1"]))
    if "Πρόσωπο2" in wb.sheetnames:
        fields.update(_read_horizontal(wb["Πρόσωπο2"]))
    fields.update(_read_horizontal(wb["Διαμονή"]))
    fields.update(_read_horizontal(wb["Δηλώσεις"]))

    children = _read_table(wb["Τέκνα"]) if "Τέκνα" in wb.sheetnames else []
    income_rows = _read_table(wb["Εισοδήματα"]) if "Εισοδήματα" in wb.sheetnames else []
    supporting_documents = (
        _read_table(wb["Δικαιολογητικά"]) if "Δικαιολογητικά" in wb.sheetnames else []
    )

    return ParsedImport(
        metadata=metadata,
        fields=fields,
        children=children,
        income_rows=income_rows,
        supporting_documents=supporting_documents,
    )


def _parse_date(value, label: str = "Ημερομηνία") -> date:
    # datetime is a subclass of date: check it first so a date (not datetime) is returned.
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = "" if value is None else str(value).strip()
    if not text:
        raise ValidationError(f"{label}: λείπει ημερομηνία.")
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        raise ValidationError(
            f"{label}: μη έγκυρη ημερομηνία «{text}» (αναμενόμενη μορφή YYYY-MM-DD)."
        )


def _parse_decimal(value, label: str = "Ποσό") -> Decimal:
    """Blank -> 0; unparsable or negative text is a validation error (never silently 0)."""
    if value is None or str(value).strip() == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        number = value
    else:
        text = str(value).replace("€", "").replace(" ", "").replace(" ", "").strip()
        if "," in text and "." in text:
            # Both separators: the right-most one is the decimal separator.
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise ValidationError(f"{label}: μη έγκυρο ποσό «{value}».")
    if not number.is_finite() or number < 0:
        raise ValidationError(f"{label}: μη έγκυρο ποσό «{value}».")
    return number
