from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from applications.import_template import (
    APPLICATION_FIELDS,
    CHILDREN_COLUMNS,
    DEFAULT_SUBMISSION_CYCLE_ID,
    INCOME_COLUMNS,
    INSTRUCTIONS_GREEK,
    METADATA_FIELDS,
    OTHER_DECLARATIONS_FIELDS,
    PERSON1_FIELDS,
    PERSON2_FIELDS,
    RESIDENCE_DECLARATION_FIELDS,
    SUPPORTING_DOC_COLUMNS,
    TEMPLATE_VERSION,
)


def _write_key_value_sheet(ws, fields, default_values=None):
    default_values = default_values or {}
    header_font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 42
    ws.append(["Πεδίο (κωδικός)", "Ελληνική περιγραφή", "Τιμή / υπόδειξη"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for key, label, hint in fields:
        value = default_values.get(key, "")
        ws.append([key, label, value])


def _write_horizontal_sheet(ws, fields):
    header_font = Font(bold=True)
    key_font = Font(italic=True, color="444444")
    for col_idx, (key, label, hint) in enumerate(fields, start=1):
        col = get_column_letter(col_idx)
        ws.column_dimensions[col].width = max(14, min(32, len(label) * 0.45 + 8))
        c1 = ws.cell(row=1, column=col_idx, value=label)
        c1.font = header_font
        c1.fill = PatternFill("solid", fgColor="D9E1F2")
        c1.alignment = Alignment(wrap_text=True, vertical="top")
        c2 = ws.cell(row=2, column=col_idx, value=key)
        c2.font = key_font
        if hint:
            c3 = ws.cell(row=3, column=col_idx, value=hint)
            c3.font = Font(color="666666", size=9)
            c3.alignment = Alignment(wrap_text=True, vertical="top")


def _write_table_sheet(ws, columns):
    header_font = Font(bold=True)
    key_font = Font(italic=True, color="444444")
    for col_idx, (key, label, hint) in enumerate(columns, start=1):
        col = get_column_letter(col_idx)
        ws.column_dimensions[col].width = max(12, min(28, len(label) * 0.4 + 6))
        ws.cell(row=1, column=col_idx, value=label).font = header_font
        ws.cell(row=1, column=col_idx).fill = PatternFill("solid", fgColor="D9E1F2")
        ws.cell(row=2, column=col_idx, value=key).font = key_font
        if hint:
            ws.cell(row=3, column=col_idx, value=hint).font = Font(color="666666", size=9)


def build_workbook():
    wb = Workbook()
    ws_instructions = wb.active
    ws_instructions.title = "Οδηγίες"
    ws_instructions.column_dimensions["A"].width = 100
    for line in INSTRUCTIONS_GREEK.format(version=TEMPLATE_VERSION).splitlines():
        ws_instructions.append([line])

    ws_meta = wb.create_sheet("Μεταδεδομένα")
    _write_key_value_sheet(
        ws_meta,
        METADATA_FIELDS,
        default_values={
            "template_version": TEMPLATE_VERSION,
            "submission_cycle_id": str(DEFAULT_SUBMISSION_CYCLE_ID),
        },
    )

    ws_app = wb.create_sheet("Αίτηση")
    _write_horizontal_sheet(ws_app, APPLICATION_FIELDS)

    ws_p1 = wb.create_sheet("Πρόσωπο1")
    _write_horizontal_sheet(ws_p1, PERSON1_FIELDS)

    ws_p2 = wb.create_sheet("Πρόσωπο2")
    _write_horizontal_sheet(ws_p2, PERSON2_FIELDS)

    ws_res = wb.create_sheet("Διαμονή")
    _write_horizontal_sheet(ws_res, RESIDENCE_DECLARATION_FIELDS)

    ws_children = wb.create_sheet("Τέκνα")
    _write_table_sheet(ws_children, CHILDREN_COLUMNS)

    ws_income = wb.create_sheet("Εισοδήματα")
    _write_table_sheet(ws_income, INCOME_COLUMNS)

    ws_decl = wb.create_sheet("Δηλώσεις")
    _write_horizontal_sheet(ws_decl, OTHER_DECLARATIONS_FIELDS)

    ws_docs = wb.create_sheet("Δικαιολογητικά")
    _write_table_sheet(ws_docs, SUPPORTING_DOC_COLUMNS)

    return wb


class Command(BaseCommand):
    help = "Generate empty Excel template for Appendix 3 application import (§9.2)."

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            type=str,
            default="",
            help="Output .xlsx path (default: static/templates/application_import_appendix3_vX.xlsx)",
        )

    def handle(self, *args, **options):
        out = options["output"]
        if not out:
            out = (
                Path(settings.BASE_DIR)
                / "static"
                / "templates"
                / f"application_import_appendix3_v{TEMPLATE_VERSION}.xlsx"
            )
        else:
            out = Path(out)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = build_workbook()
        wb.save(out)
        self.stdout.write(self.style.SUCCESS(f"Wrote {out}"))
